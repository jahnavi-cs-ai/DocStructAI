import io
from openpyxl import load_workbook
from docustruct.pipeline import ExtractedTable, export_excel

def test_excel_export_sheets():
    table = ExtractedTable(["ID", "Amount"], [["1", "10"], ["2", "20"]], [1], name="Table_01")
    data = export_excel([table], [], {"pages": 1})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Table_01", "Validation", "Provenance"]
    assert wb["Table_01"]["B3"].value == "20"
